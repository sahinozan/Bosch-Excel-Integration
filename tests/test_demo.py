from pathlib import Path
import sys

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Code"))
from helper import convert_schedule, transform


def test_component_totals_and_missing_products():
    schedule = pd.DataFrame([
        ["2020-01-06", "SYN-LINE-A", "Morning", "SYN-PRODUCT-A", 10],
        ["2020-01-06", "SYN-LINE-A", "Morning", "SYN-PRODUCT-A", 5],
        ["2020-01-06", "SYN-LINE-A", "Morning", "SYN-UNMAPPED", 7],
    ], columns=["date", "line", "shift", "product_id", "quantity"])
    bom = pd.DataFrame([
        ["SYN-PRODUCT-A", "SYN-PART-A", 2], ["SYN-PRODUCT-A", "SYN-PART-B", 3],
    ], columns=["product_id", "component_id", "units_per_product"])
    _, totals, missing = convert_schedule(schedule, bom)
    assert totals.set_index("component_id")["component_quantity"].to_dict() == {
        "SYN-PART-A": 30, "SYN-PART-B": 45,
    }
    assert missing["quantity"].sum() == 7
    with pytest.raises(ValueError, match="Duplicate"):
        convert_schedule(schedule, pd.concat([bom, bom]))


def test_workbook_export(tmp_path):
    path = tmp_path / "plan.xlsx"
    detail, totals, missing = transform(ROOT / "examples", path)
    assert len(detail) == 60 and len(totals) == 60 and missing.empty
    book = load_workbook(path)
    assert book.sheetnames == ["Detail", "Totals", "Unmapped"]
    assert book.properties.creator == book.properties.lastModifiedBy == "Public demo"
    assert not book._external_links
    assert all(sheet.sheet_state == "visible" for sheet in book)


def test_invalid_quantity_rejected():
    schedule = pd.read_csv(ROOT / "examples" / "schedule.csv")
    bom = pd.read_csv(ROOT / "examples" / "bill_of_materials.csv")
    schedule.loc[0, "quantity"] = -1
    with pytest.raises(ValueError, match="nonnegative"):
        convert_schedule(schedule, bom)
